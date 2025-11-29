import os
from openpyxl import Workbook, load_workbook

def log_to_excel(file_path, test_name, flow, status, exec_time, error_message):
    """
    Append test result to Excel file (creates it if not exists)
    """
    if not os.path.exists(file_path):
        # Create new workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        ws.append(["Test Name", "Flow", "Status", "Execution Time (s)", "Error Message"])
        wb.save(file_path)

    # Load existing file
    wb = load_workbook(file_path)
    ws = wb.active

    # Append result
    ws.append([test_name, flow, status, exec_time, error_message])
    wb.save(file_path)
